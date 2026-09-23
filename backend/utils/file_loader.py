from __future__ import annotations

import os
from pathlib import Path

import fitz
from docx import Document


def extract_text_from_pdf(file_path: str) -> str:
    try:
        with fitz.open(file_path) as document:
            return "\n".join(page.get_text() for page in document).strip()
    except Exception as exc:
        raise ValueError(f"Failed to extract PDF text: {exc}") from exc


def extract_text_from_docx(file_path: str) -> str:
    try:
        document = Document(file_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ValueError(f"Failed to extract DOCX text: {exc}") from exc


def extract_text_from_txt(file_path: str) -> str:
    try:
        return Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
    except Exception as exc:
        raise ValueError(f"Failed to extract TXT text: {exc}") from exc


def extract_resume_text(file_path: str) -> str:
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".pdf":
        return extract_text_from_pdf(file_path)
    if extension == ".docx":
        return extract_text_from_docx(file_path)
    if extension == ".txt":
        return extract_text_from_txt(file_path)
    raise ValueError("Unsupported file type. Only PDF, DOCX, and TXT are supported.")
